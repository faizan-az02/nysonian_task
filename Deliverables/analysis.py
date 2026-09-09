import csv
import math
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path
from statistics import median
from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

WORKBOOK = Path("Supply Chain Assessment Data set.xlsx")

BUDGET = 500_000
CONTAINER_LIMIT = 2
MOQ_DEFAULT = 24

HORIZONS = (30, 60, 90)

WAREHOUSE_ALIASES = {
    "Warehouse-A-legacy": "Warehouse A"
}


# ============================================================
# BASIC HELPERS
# ============================================================

def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)[:10]).date()


def percentile(values, pct):
    values = sorted(values)

    position = (len(values) - 1) * pct
    low = math.floor(position)
    high = math.ceil(position)

    if low == high:
        return values[low]

    return values[low] + (
        values[high] - values[low]
    ) * (position - low)


# ============================================================
# 1. LOAD ALL EXCEL SHEETS
# ============================================================

wb = load_workbook(
    WORKBOOK,
    data_only=True,
    read_only=True
)

tables = {}

for ws in wb.worksheets:

    rows = list(
        ws.iter_rows(values_only=True)
    )

    if not rows:
        tables[ws.title] = []
        continue

    headers = [
        str(value).strip()
        if value is not None
        else f"Column_{i}"
        for i, value in enumerate(
            rows[0],
            start=1
        )
    ]

    tables[ws.title] = [
        dict(zip(headers, row))
        for row in rows[1:]
        if any(
            value not in (None, "")
            for value in row
        )
    ]


products = {
    row["sku"]: row
    for row in tables["Products"]
}

active_skus = {
    sku
    for sku, product in products.items()
    if product["status"] == "Active"
}

discontinued_skus = {
    sku
    for sku, product in products.items()
    if product["status"] == "Discontinued"
}

warehouses = sorted({
    WAREHOUSE_ALIASES.get(
        row["warehouse"],
        row["warehouse"]
    )
    for row in tables["Sales Daily"]
})

sales_dates = [
    as_date(row["date"])
    for row in tables["Sales Daily"]
]

as_of = max(sales_dates)
min_date = min(sales_dates)


# ============================================================
# 2. PROMOTION LIFT
# ============================================================

promo_lift = {}
category_promos = defaultdict(list)

for row in tables["Promotions"]:

    sku = row.get("sku")
    lift = row.get("expected_lift_pct")

    if (
        sku in products
        and isinstance(lift, (int, float))
    ):

        promo_lift[
            (row["promotion_id"], sku)
        ] = lift

        category_promos[
            products[sku]["category"]
        ].append(lift)


category_lift = {
    category: median(values)
    for category, values
    in category_promos.items()
}

all_promo_lifts = [
    value
    for values in category_promos.values()
    for value in values
]

global_lift = median(all_promo_lifts)


# ============================================================
# 3. CLEAN / AGGREGATE SALES
# ============================================================

daily_sales = defaultdict(
    lambda: {
        "units": 0.0,
        "stockout": False
    }
)

for row in tables["Sales Daily"]:

    sku = row["sku"]

    warehouse = WAREHOUSE_ALIASES.get(
        row["warehouse"],
        row["warehouse"]
    )

    lift = 0

    if row.get("promotion_id"):

        lift = promo_lift.get(
            (
                row["promotion_id"],
                sku
            ),
            category_lift.get(
                products[sku]["category"],
                global_lift
            )
        )

    key = (
        sku,
        warehouse,
        as_date(row["date"])
    )

    daily_sales[key]["units"] += (
        row["units"] / (1 + lift)
    )

    if row.get("stockout_flag") == "Y":
        daily_sales[key]["stockout"] = True


# ============================================================
# 4. INVENTORY LOOKUP
# ============================================================

inventory = {}

for row in tables["Inventory Daily"]:

    warehouse = WAREHOUSE_ALIASES.get(
        row["warehouse"],
        row["warehouse"]
    )

    inventory[
        (
            row["sku"],
            warehouse,
            as_date(row["date"])
        )
    ] = row


# ============================================================
# 5. OPEN PURCHASE ORDERS
# ============================================================

open_po = defaultdict(float)

for row in tables["Purchase Orders"]:

    if row["status"] not in {
        "Open",
        "Partially Received"
    }:
        continue

    remaining = max(
        0,
        (row["ordered_qty"] or 0)
        - (row["received_qty"] or 0)
    )

    if (
        as_date(row["stated_eta"])
        <= as_of + timedelta(days=90)
    ):

        warehouse = WAREHOUSE_ALIASES.get(
            row["warehouse"],
            row["warehouse"]
        )

        open_po[
            (
                row["sku"],
                warehouse
            )
        ] += remaining


# ============================================================
# 6. SUPPLIER LEAD TIMES
# ============================================================

supplier_leads = defaultdict(list)
supplier_slips = defaultdict(list)

for row in tables["Purchase Orders"]:

    if not row.get("actual_receipt_date"):
        continue

    supplier = row["supplier"]

    actual = as_date(
        row["actual_receipt_date"]
    )

    ordered = as_date(
        row["order_date"]
    )

    eta = as_date(
        row["stated_eta"]
    )

    supplier_leads[supplier].append(
        (actual - ordered).days
    )

    supplier_slips[supplier].append(
        (actual - eta).days
    )


lead_stats = {}

for supplier, values in supplier_leads.items():

    lead_stats[supplier] = {
        "median": median(values),
        "p75": percentile(values, 0.75),
        "p90": percentile(values, 0.90),
        "median_eta_slip":
            median(
                supplier_slips[supplier]
            )
    }


# ============================================================
# 7. MOQ FROM HISTORICAL PURCHASE ORDERS
# ============================================================

moq_values = defaultdict(list)

for row in tables["Purchase Orders"]:

    moq_values[
        row["sku"]
    ].append(
        row["ordered_qty"]
    )


moq = {
    sku: min(values)
    for sku, values
    in moq_values.items()
}


# ============================================================
# CORE FORECAST HELPERS
# ============================================================

def is_stockout(sku, warehouse, day):

    sales = daily_sales.get(
        (sku, warehouse, day)
    )

    inv = inventory.get(
        (sku, warehouse, day)
    )

    return (
        bool(
            sales
            and sales["stockout"]
        )
        or bool(
            inv
            and isinstance(
                inv.get("available"),
                (int, float)
            )
            and inv["available"] <= 0
        )
    )


@lru_cache(maxsize=None)
def clean_values(
    sku,
    warehouse,
    cutoff,
    days_back,
    weekday=None
):

    start = max(
        min_date,
        cutoff - timedelta(
            days=days_back - 1
        ),
        as_date(
            products[sku]["launch_date"]
        )
    )

    values = []

    day = start

    while day <= cutoff:

        if (
            weekday is None
            or day.weekday() == weekday
        ):

            if not is_stockout(
                sku,
                warehouse,
                day
            ):

                values.append(
                    daily_sales.get(
                        (
                            sku,
                            warehouse,
                            day
                        ),
                        {"units": 0}
                    )["units"]
                )

        day += timedelta(days=1)

    return tuple(values)


@lru_cache(maxsize=None)
def fallback_rates(cutoff):

    category_rates = defaultdict(list)
    sku_rates = defaultdict(list)

    for sku in active_skus:

        if (
            as_date(
                products[sku]["launch_date"]
            )
            > cutoff
        ):
            continue

        for warehouse in warehouses:

            values = clean_values(
                sku,
                warehouse,
                cutoff,
                180
            )

            if not values:
                continue

            rate = (
                sum(values) / len(values)
            )

            category_rates[
                (
                    products[sku]["category"],
                    warehouse
                )
            ].append(rate)

            sku_rates[sku].append(rate)


    category_fallback = {
        key: median(values)
        for key, values
        in category_rates.items()
    }

    sku_fallback = {
        sku:
            sum(values) / len(values)
        for sku, values
        in sku_rates.items()
    }

    all_rates = [
        value
        for values in sku_rates.values()
        for value in values
    ]

    global_rate = (
        sum(all_rates) / len(all_rates)
        if all_rates
        else 0
    )

    return (
        category_fallback,
        sku_fallback,
        global_rate
    )


def baseline_rate(
    sku,
    warehouse,
    cutoff
):

    if (
        sku in discontinued_skus
        or as_date(
            products[sku]["launch_date"]
        ) > cutoff
    ):
        return 0

    values = clean_values(
        sku,
        warehouse,
        cutoff,
        180
    )

    if len(values) >= 30:
        return sum(values) / len(values)

    category_rates, sku_rates, global_rate = (
        fallback_rates(cutoff)
    )

    return category_rates.get(
        (
            products[sku]["category"],
            warehouse
        ),
        sku_rates.get(
            sku,
            global_rate
        )
    )


def forecast(
    sku,
    warehouse,
    cutoff,
    horizon
):

    if (
        sku in discontinued_skus
        or as_date(
            products[sku]["launch_date"]
        ) > cutoff
    ):
        return 0

    category_rates, sku_rates, global_rate = (
        fallback_rates(cutoff)
    )

    default = category_rates.get(
        (
            products[sku]["category"],
            warehouse
        ),
        sku_rates.get(
            sku,
            global_rate
        )
    )

    rates = {}

    for days in (28, 90, 180):

        values = clean_values(
            sku,
            warehouse,
            cutoff,
            days
        )

        rates[days] = (
            sum(values) / len(values)
            if values
            else None
        )


    weekday_rates = {}

    for weekday in range(7):

        values = clean_values(
            sku,
            warehouse,
            cutoff,
            180,
            weekday
        )

        weekday_rates[weekday] = (
            sum(values) / len(values)
            if values
            else None
        )


    total = 0

    for offset in range(
        1,
        horizon + 1
    ):

        day = (
            cutoff
            + timedelta(days=offset)
        )

        pieces = [
            (0.45, rates[28]),
            (0.25, rates[90]),
            (0.15, rates[180]),
            (
                0.15,
                weekday_rates[
                    day.weekday()
                ]
            )
        ]

        weighted_sum = sum(
            weight * value
            for weight, value in pieces
            if value is not None
        )

        weight_sum = sum(
            weight
            for weight, value in pieces
            if value is not None
        )

        total += (
            weighted_sum / weight_sum
            if weight_sum
            else default
        )

    return total


def latest_available(
    sku,
    warehouse,
    cutoff
):

    dates = [
        day
        for product_sku, wh, day
        in inventory
        if (
            product_sku == sku
            and wh == warehouse
            and day <= cutoff
        )
    ]

    if not dates:
        return 0

    value = inventory[
        (
            sku,
            warehouse,
            max(dates)
        )
    ].get("available")

    return (
        max(0, value)
        if isinstance(
            value,
            (int, float)
        )
        else 0
    )


# ============================================================
# 8. MODEL VALIDATION
# ============================================================

validation_rows = []

validation_cutoffs = (
    date(2025, 10, 1),
    date(2026, 1, 1),
    date(2026, 4, 1)
)


for cutoff in validation_cutoffs:

    eligible_skus = [
        sku
        for sku in active_skus
        if as_date(
            products[sku]["launch_date"]
        ) <= cutoff
    ]

    for horizon in HORIZONS:

        for model in (
            "baseline",
            "challenger"
        ):

            actual_total = 0
            abs_error = 0
            signed_error = 0

            dollar_actual = 0
            dollar_error = 0

            under = 0
            over = 0

            actual_risks = 0
            predicted_risks = 0
            true_positive = 0

            censored_days = 0


            for sku in eligible_skus:

                cost = products[sku][
                    "unit_cost"
                ]

                for warehouse in warehouses:

                    # -------------------------
                    # Actual unconstrained demand
                    # -------------------------

                    impute_rate = baseline_rate(
                        sku,
                        warehouse,
                        cutoff
                    )

                    actual = 0

                    for offset in range(
                        1,
                        horizon + 1
                    ):

                        day = (
                            cutoff
                            + timedelta(
                                days=offset
                            )
                        )

                        observed = daily_sales.get(
                            (
                                sku,
                                warehouse,
                                day
                            ),
                            {"units": 0}
                        )["units"]

                        if is_stockout(
                            sku,
                            warehouse,
                            day
                        ):

                            censored_days += 1

                            actual += max(
                                observed,
                                impute_rate
                            )

                        else:
                            actual += observed


                    # -------------------------
                    # Forecast
                    # -------------------------

                    if model == "baseline":

                        predicted = (
                            baseline_rate(
                                sku,
                                warehouse,
                                cutoff
                            )
                            * horizon
                        )

                    else:

                        predicted = forecast(
                            sku,
                            warehouse,
                            cutoff,
                            horizon
                        )


                    error = (
                        predicted - actual
                    )

                    available = latest_available(
                        sku,
                        warehouse,
                        cutoff
                    )

                    actual_risk = (
                        available < actual
                    )

                    predicted_risk = (
                        available < predicted
                    )


                    actual_total += actual

                    abs_error += abs(error)

                    signed_error += error

                    dollar_actual += (
                        actual * cost
                    )

                    dollar_error += (
                        abs(error) * cost
                    )

                    under += max(
                        actual - predicted,
                        0
                    )

                    over += max(
                        predicted - actual,
                        0
                    )

                    actual_risks += int(
                        actual_risk
                    )

                    predicted_risks += int(
                        predicted_risk
                    )

                    true_positive += int(
                        actual_risk
                        and predicted_risk
                    )


            validation_rows.append({

                "cutoff":
                    cutoff.isoformat(),

                "model":
                    model,

                "horizon_days":
                    horizon,

                "wmape":
                    abs_error
                    / actual_total,

                "bias_pct":
                    signed_error
                    / actual_total,

                "dollar_weighted_mape":
                    dollar_error
                    / dollar_actual,

                "under_forecast_units":
                    under,

                "over_forecast_units":
                    over,

                "service_risk_recall":
                    (
                        true_positive
                        / actual_risks
                        if actual_risks
                        else 1
                    ),

                "predicted_risk_count":
                    predicted_risks,

                "actual_risk_count":
                    actual_risks,

                "validation_censored_days":
                    censored_days
            })


# ============================================================
# 9. VALIDATION ERROR FOR FORECAST INTERVALS
# ============================================================

error_by_horizon = defaultdict(list)

for row in validation_rows:

    if row["model"] == "challenger":

        error_by_horizon[
            row["horizon_days"]
        ].append(
            row["wmape"]
        )


avg_error = {
    horizon:
        (
            sum(
                error_by_horizon[
                    horizon
                ]
            )
            / len(
                error_by_horizon[
                    horizon
                ]
            )
            if error_by_horizon[
                horizon
            ]
            else 0.30
        )
    for horizon in HORIZONS
}


# ============================================================
# 10. CREATE REORDER CANDIDATES
# ============================================================

candidates = []


for sku in active_skus:

    product = products[sku]

    supplier = product["supplier"]

    lead_days = int(
        round(
            (
                lead_stats.get(
                    supplier,
                    {}
                ).get("p75")
                or 60
            )
        )
    )

    target_days = min(
        90,
        max(
            30,
            lead_days
        ) + 7
    )

    sku_moq = moq.get(
        sku,
        MOQ_DEFAULT
    )


    for warehouse in warehouses:

        forecasts = {
            horizon:
                forecast(
                    sku,
                    warehouse,
                    as_of,
                    horizon
                )
            for horizon
            in HORIZONS
        }

        available = latest_available(
            sku,
            warehouse,
            as_of
        )

        inbound = open_po[
            (
                sku,
                warehouse
            )
        ]

        position = (
            available + inbound
        )

        target = forecast(
            sku,
            warehouse,
            as_of,
            target_days
        )

        raw_need = max(
            0,
            target - position
        )

        candidate_type = "URGENT_SHORTAGE"
        reason_code = "BUY_APPROVAL_REQUIRED"

        if raw_need <= 0:
            target_days = 180
            target = forecast(
                sku,
                warehouse,
                as_of,
                target_days
            )
            raw_need = max(
                0,
                target - position
            )
            candidate_type = "BUDGET_CAPACITY_FILL"
            reason_code = "OPTIONAL_BUY_USES_REMAINING_BUDGET"

        recommended_qty = int(
            math.ceil(
                raw_need / sku_moq
            )
            * sku_moq
        )

        if recommended_qty <= 0:
            continue


        shortage_30 = max(
            0,
            forecasts[30] - position
        )

        shortage_60 = max(
            0,
            forecasts[60] - position
        )

        shortage_90 = max(
            0,
            forecasts[90] - position
        )


        margin = (
            product["list_price"]
            - product["unit_cost"]
        )


        priority = (
            shortage_30 * 3
            + shortage_60 * 2
            + shortage_90
            + shortage_90
            * margin / 100
        )

        if candidate_type == "BUDGET_CAPACITY_FILL":
            priority = (
                forecasts[90]
                * max(margin, 0)
                / 100
            )


        candidates.append({

            "supplier":
                supplier,

            "sku":
                sku,

            "product_name":
                product["product_name"],

            "category":
                product["category"],

            "warehouse":
                warehouse,

            "available":
                available,

            "open_po_units_90d":
                inbound,

            "inventory_position":
                position,

            "forecast_30":
                forecasts[30],

            "forecast_60":
                forecasts[60],

            "forecast_90":
                forecasts[90],

            "lead_cover_days":
                max(
                    30,
                    min(
                        90,
                        lead_days
                    )
                ),

            "target_days":
                target_days,

            "moq_proxy":
                sku_moq,

            "recommended_qty":
                recommended_qty,

            "candidate_type":
                candidate_type,

            "reason_code":
                reason_code,

            "unit_cost":
                product["unit_cost"],

            "purchase_value":
                recommended_qty
                * product["unit_cost"],

            "shortage_30":
                shortage_30,

            "shortage_60":
                shortage_60,

            "shortage_90":
                shortage_90,

            "priority_score":
                priority
        })


# ============================================================
# 11. SELECT TOP 2 SUPPLIERS / CONTAINERS
# ============================================================

supplier_scores = defaultdict(float)

score_pool = [
    row
    for row in candidates
    if row["candidate_type"]
    == "URGENT_SHORTAGE"
]

if not score_pool:
    score_pool = candidates

for row in score_pool:

    supplier_scores[
        row["supplier"]
    ] += row["priority_score"]


selected_suppliers = [
    supplier
    for supplier, score
    in sorted(
        supplier_scores.items(),
        key=lambda x: x[1],
        reverse=True
    )[:CONTAINER_LIMIT]
]


container_slots = {
    supplier: i + 1
    for i, supplier
    in enumerate(
        selected_suppliers
    )
}


remaining_budget = BUDGET

selected = []


eligible_candidates = [
    row
    for row in candidates
    if row["supplier"]
    in selected_suppliers
]


eligible_candidates.sort(
    key=lambda row:
        (
            row["candidate_type"]
            == "URGENT_SHORTAGE",
            row["priority_score"]
        ),
    reverse=True
)


for row in eligible_candidates:

    if (
        row["purchase_value"]
        <= remaining_budget
    ):

        row["container_slot"] = (
            container_slots[
                row["supplier"]
            ]
        )

        selected.append(row)

        remaining_budget -= (
            row["purchase_value"]
        )


selected_lookup = {
    (
        row["sku"],
        row["warehouse"]
    ): row
    for row in selected
}


# ============================================================
# 12. BUILD MAIN FORECAST OUTPUT
# ============================================================

forecast_rows = []
required_rows = []


for sku in sorted(active_skus):

    product = products[sku]

    sku_totals = defaultdict(float)


    for warehouse in warehouses:

        forecasts = {
            horizon:
                forecast(
                    sku,
                    warehouse,
                    as_of,
                    horizon
                )
            for horizon
            in HORIZONS
        }


        for horizon in HORIZONS:

            sku_totals[
                horizon
            ] += forecasts[horizon]


        available = latest_available(
            sku,
            warehouse,
            as_of
        )

        inbound = open_po[
            (
                sku,
                warehouse
            )
        ]

        position = (
            available + inbound
        )


        reorder = selected_lookup.get(
            (
                sku,
                warehouse
            ),
            {}
        )


        if position < forecasts[30]:
            risk = "Critical"

        elif position < forecasts[60]:
            risk = "High"

        elif position < forecasts[90]:
            risk = "Medium"

        else:
            risk = "Low"


        for horizon in HORIZONS:

            point_forecast = (
                forecasts[horizon]
            )

            error = (
                avg_error[horizon]
            )


            required_rows.append({

                "as_of_date":
                    as_of.isoformat(),

                "sku":
                    sku,

                "product_name":
                    product[
                        "product_name"
                    ],

                "category":
                    product[
                        "category"
                    ],

                "warehouse":
                    warehouse,

                "supplier":
                    product[
                        "supplier"
                    ],

                "forecast_horizon_days":
                    horizon,

                "point_forecast_units":
                    point_forecast,

                "forecast_low_units":
                    max(
                        0,
                        point_forecast
                        * (1 - error)
                    ),

                "forecast_high_units":
                    point_forecast
                    * (1 + error),

                "available_inventory":
                    available,

                "inbound_quantity_90d":
                    inbound,

                "inventory_position":
                    position,

                "expected_lead_time_days_p75":
                    lead_stats.get(
                        product["supplier"],
                        {}
                    ).get(
                        "p75",
                        ""
                    ),

                "reorder_date":
                    (
                        as_of.isoformat()
                        if reorder
                        else ""
                    ),

                "recommended_quantity":
                    reorder.get(
                        "recommended_qty",
                        0
                    ),

                "estimated_purchase_cost":
                    reorder.get(
                        "purchase_value",
                        0
                    ),

                "priority_or_risk_classification":
                    risk,

                "reason_code":
                    reorder.get(
                        "reason_code",
                        "NO_REORDER_POSITION_SUFFICIENT"
                    ),

                "projected_stockout_date":
                    "",

                "container_slot":
                    reorder.get(
                        "container_slot",
                        ""
                    ),

                "assumptions":
                    (
                        "Gross units demand; "
                        "promo de-lifted; "
                        "stockout days censored; "
                        "MOQ from historical min PO; "
                        "two supplier/container slots."
                    ),

                "human_approval_status":
                    (
                        "Pending approval"
                        if reorder
                        else
                        "No purchase action"
                    )
            })


    forecast_rows.append({

        "sku":
            sku,

        "product_name":
            product[
                "product_name"
            ],

        "category":
            product[
                "category"
            ],

        "forecast_30":
            sku_totals[30],

        "forecast_60":
            sku_totals[60],

        "forecast_90":
            sku_totals[90]
    })


forecast_rows = sorted(
    forecast_rows,
    key=lambda row:
        row["forecast_90"],
    reverse=True
)[:10]


# ============================================================
# 13. REORDER OUTPUT
# ============================================================

reorder_rows = [
    {
        "container_slot":
            row["container_slot"],

        "supplier":
            row["supplier"],

        "sku":
            row["sku"],

        "product_name":
            row["product_name"],

        "warehouse":
            row["warehouse"],

        "available":
            row["available"],

        "open_po_units_90d":
            row["open_po_units_90d"],

        "forecast_30":
            row["forecast_30"],

        "forecast_60":
            row["forecast_60"],

        "forecast_90":
            row["forecast_90"],

        "lead_cover_days":
            row["lead_cover_days"],

        "target_days":
            row["target_days"],

        "moq_proxy":
            row["moq_proxy"],

        "recommended_qty":
            row["recommended_qty"],

        "candidate_type":
            row["candidate_type"],

        "reason_code":
            row["reason_code"],

        "unit_cost":
            row["unit_cost"],

        "purchase_value":
            row["purchase_value"],

        "shortage_30":
            row["shortage_30"],

        "shortage_60":
            row["shortage_60"],

        "shortage_90":
            row["shortage_90"]
    }

    for row in selected
]


# ============================================================
# 14. AUDIT OUTPUT
# ============================================================

sales_panel_days = (
    as_of - min_date
).days + 1

expected_sales_panel_rows = (
    len(active_skus)
    * len(warehouses)
    * sales_panel_days
)

observed_sales_panel_rows = len({
    (
        row["sku"],
        WAREHOUSE_ALIASES.get(
            row["warehouse"],
            row["warehouse"]
        ),
        as_date(row["date"])
    )
    for row in tables["Sales Daily"]
    if row["sku"] in active_skus
})

inventory_keys = [
    (
        row["sku"],
        WAREHOUSE_ALIASES.get(
            row["warehouse"],
            row["warehouse"]
        ),
        as_date(row["date"])
    )
    for row in tables["Inventory Daily"]
]

received_pos = [
    row
    for row in tables["Purchase Orders"]
    if row.get("actual_receipt_date")
]

delayed_received_pos = [
    row
    for row in received_pos
    if as_date(row["actual_receipt_date"])
    > as_date(row["stated_eta"])
]

status_counts = Counter(
    row["status"]
    for row in tables["Products"]
)

audit_rows = [
    {
        "area": "Sales",
        "finding": "Sparse order-line grain, not complete daily demand panel",
        "count": expected_sales_panel_rows - observed_sales_panel_rows,
        "pct_affected": (
            (expected_sales_panel_rows - observed_sales_panel_rows)
            / expected_sales_panel_rows
        ),
        "decision_impact": "Complete the SKU-warehouse-day calendar before forecasting"
    },
    {
        "area": "Sales",
        "finding": "Stockout sales are censored",
        "count": sum(
            1
            for row in tables["Sales Daily"]
            if row.get("stockout_flag") == "Y"
        ),
        "pct_affected": (
            sum(
                1
                for row in tables["Sales Daily"]
                if row.get("stockout_flag") == "Y"
            )
            / len(tables["Sales Daily"])
        ),
        "decision_impact": "Impute demand on stockout days; do not treat low sales as low demand"
    },
    {
        "area": "Sales",
        "finding": "Return rows exist and should not be netted from gross demand by default",
        "count": sum(
            1
            for row in tables["Sales Daily"]
            if (row.get("return_units") or 0) > 0
        ),
        "pct_affected": (
            sum(
                1
                for row in tables["Sales Daily"]
                if (row.get("return_units") or 0) > 0
            )
            / len(tables["Sales Daily"])
        ),
        "decision_impact": "Keep returns as an exception/quality feature unless pre-fulfillment cancellation"
    },
    {
        "area": "Inventory",
        "finding": "Negative balances found in on_hand, allocated, or available",
        "count": sum(
            1
            for row in tables["Inventory Daily"]
            if any(
                isinstance(row.get(column), (int, float))
                and row.get(column) < 0
                for column in ("on_hand", "allocated", "available")
            )
        ),
        "pct_affected": (
            sum(
                1
                for row in tables["Inventory Daily"]
                if any(
                    isinstance(row.get(column), (int, float))
                    and row.get(column) < 0
                    for column in ("on_hand", "allocated", "available")
                )
            )
            / len(tables["Inventory Daily"])
        ),
        "decision_impact": "Flag inventory confidence before recommending purchase action"
    },
    {
        "area": "Inventory",
        "finding": "Duplicate SKU-warehouse-date inventory rows",
        "count": sum(
            count - 1
            for count in Counter(inventory_keys).values()
            if count > 1
        ),
        "pct_affected": (
            sum(
                count - 1
                for count in Counter(inventory_keys).values()
                if count > 1
            )
            / len(tables["Inventory Daily"])
        ),
        "decision_impact": "Deduplicate latest snapshot before calculating inventory position"
    },
    {
        "area": "Inventory",
        "finding": "Legacy warehouse naming appears",
        "count": sum(
            1
            for row in tables["Inventory Daily"] + tables["Sales Daily"]
            if row.get("warehouse") in WAREHOUSE_ALIASES
        ),
        "pct_affected": (
            sum(
                1
                for row in tables["Inventory Daily"] + tables["Sales Daily"]
                if row.get("warehouse") in WAREHOUSE_ALIASES
            )
            / (
                len(tables["Inventory Daily"])
                + len(tables["Sales Daily"])
            )
        ),
        "decision_impact": "Normalize warehouse names before joins and reporting"
    },
    {
        "area": "POs",
        "finding": "Received POs arrived after stated ETA",
        "count": len(delayed_received_pos),
        "pct_affected": (
            len(delayed_received_pos)
            / len(received_pos)
            if received_pos
            else 0
        ),
        "decision_impact": "Use actual supplier p75 lead time rather than stated ETA only"
    },
    {
        "area": "POs",
        "finding": "Open or partially received POs require inbound netting",
        "count": sum(
            1
            for row in tables["Purchase Orders"]
            if row["status"] in {"Open", "Partially Received"}
        ),
        "pct_affected": (
            sum(
                1
                for row in tables["Purchase Orders"]
                if row["status"] in {"Open", "Partially Received"}
            )
            / len(tables["Purchase Orders"])
        ),
        "decision_impact": "Net remaining open PO quantity before recommending new buys"
    },
    {
        "area": "Promotions",
        "finding": "Promotion lift estimates are missing",
        "count": sum(
            1
            for row in tables["Promotions"]
            if row.get("expected_lift_pct") in (None, "")
        ),
        "pct_affected": (
            sum(
                1
                for row in tables["Promotions"]
                if row.get("expected_lift_pct") in (None, "")
            )
            / len(tables["Promotions"])
        ),
        "decision_impact": "Impute category lift for modeling; require planner review for promo buys"
    },
    {
        "area": "Products",
        "finding": "Active, discontinued, and planned SKUs are mixed",
        "count": len(tables["Products"]),
        "pct_affected": 1,
        "decision_impact": (
            f"Use lifecycle controls: "
            f"{status_counts.get('Active', 0)} active, "
            f"{status_counts.get('Discontinued', 0)} discontinued, "
            f"{status_counts.get('Planned', 0)} planned"
        )
    }
]


# ============================================================
# 15. SAVE CSV FILES
# ============================================================

outputs = {
    "model_validation_results.csv": validation_rows,
    "selected_sku_forecasts.csv": forecast_rows,
    "reorder_recommendations.csv": reorder_rows,
    "sku_warehouse_forecast_recommendation.csv": required_rows,
    "Data Audit.csv": audit_rows
}

for filename, rows in outputs.items():

    if not rows:
        continue

    with open(
        filename,
        "w",
        newline="",
        encoding="utf-8"
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=list(rows[0].keys())
        )

        writer.writeheader()

        for row in rows:
            writer.writerow({
                key: round(value, 4)
                if isinstance(value, float)
                else value
                for key, value in row.items()
            })


# ============================================================
# DONE
# ============================================================

print("Outputs refreshed.")
print("Validation rows:", len(validation_rows))
print("Main output rows:", len(required_rows))
print("Pending approval reorder lines:", len(selected))
